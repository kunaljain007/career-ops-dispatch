import os
import sys
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

sys.stdout.reconfigure(encoding='utf-8')

print("Starting full generation of World-Class Sovereign Light Theme Presentation...")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# ==============================================================================
# WORLD-CLASS SOVEREIGN DEFENCE LIGHT COLOR PALETTE
# ==============================================================================
BG_CANVAS      = RGBColor(248, 250, 252)   # #F8FAFC Ultra-clean architectural ivory/slate
PANEL_WHITE    = RGBColor(255, 255, 255)   # #FFFFFF Pure crisp card surface
BORDER_SUBTLE  = RGBColor(226, 232, 240)   # #E2E8F0 Structural light slate border
BORDER_NAVY    = RGBColor(15, 30, 54)      # #0F1E36 Deep Imperial Navy border
NAVY_PRIMARY   = RGBColor(11, 27, 51)      # #0B1B33 Deep Sovereign Navy (Headlines)
NAVY_SECONDARY = RGBColor(24, 43, 73)      # #182B49 Medium Imperial Navy
GOLD_SOVEREIGN = RGBColor(176, 125, 20)    # #B07D14 Rich Atmanirbhar Gold
GOLD_BG_TINT   = RGBColor(254, 250, 240)   # #FEFAF0 Soft Gold Callout background
GOLD_BORDER    = RGBColor(217, 164, 65)    # #D9A441 Gold Accent Border
TEXT_DARK      = RGBColor(15, 23, 42)      # #0F172A Primary dark text
TEXT_MUTED     = RGBColor(64, 78, 98)      # #404E62 High-readability body text
TEXT_DIM       = RGBColor(130, 145, 165)   # #8291A5 Technical metadata
TAG_BG_LIGHT   = RGBColor(241, 245, 249)   # #F1F5F9 Tag badge background
ARMY_OLIVE     = RGBColor(45, 74, 34)      # #2D4A22 Tactical Military Green
CRIMSON_RED    = RGBColor(185, 28, 28)     # #B91C1C High-priority alert red

FONT_HEAD = "Segoe UI"
FONT_BODY = "Segoe UI"
FONT_TECH = "Consolas"

blank_layout = prs.slide_layouts[6]
IMAGE_PATH = r"C:\Users\My\.gemini\antigravity\scratch\career-ops\heatbolt_national_traction.png"

def base_slide(kicker, title, slide_idx=""):
    s = prs.slides.add_slide(blank_layout)
    
    # Background Canvas
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_CANVAS
    bg.line.fill.background()
    
    # Top Sovereign Tri-Color Inspired Accent Stripe (Navy + Gold)
    bar_navy = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.06))
    bar_navy.fill.solid()
    bar_navy.fill.fore_color.rgb = NAVY_PRIMARY
    bar_navy.line.fill.background()
    
    bar_gold = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.06), Inches(2.2), Inches(0.035))
    bar_gold.fill.solid()
    bar_gold.fill.fore_color.rgb = GOLD_SOVEREIGN
    bar_gold.line.fill.background()
    
    # Kicker
    tb = s.shapes.add_textbox(Inches(0.8), Inches(0.38), Inches(9.5), Inches(0.28))
    tf = tb.text_frame
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = f"//  {kicker.upper()}"
    p.font.name = FONT_TECH
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = GOLD_SOVEREIGN
    
    # Title
    tb_t = s.shapes.add_textbox(Inches(0.8), Inches(0.66), Inches(11.0), Inches(0.65))
    tf_t = tb_t.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
    p2 = tf_t.paragraphs[0]
    p2.text = title
    p2.font.name = FONT_HEAD
    p2.font.size = Pt(21)
    p2.font.bold = True
    p2.font.color.rgb = NAVY_PRIMARY
    
    # Gold divider rule
    rule = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.36), Inches(1.8), Inches(0.04))
    rule.fill.solid()
    rule.fill.fore_color.rgb = GOLD_SOVEREIGN
    rule.line.fill.background()
    
    # Top Right Indicator
    tb_r = s.shapes.add_textbox(Inches(8.5), Inches(0.38), Inches(4.033), Inches(0.28))
    tf_r = tb_r.text_frame
    tf_r.margin_left = tf_r.margin_top = tf_r.margin_right = tf_r.margin_bottom = 0
    pr = tf_r.paragraphs[0]
    pr.alignment = PP_ALIGN.RIGHT
    pr.text = f"DEFENCE STRATEGIC BRIEFING  |  {slide_idx}"
    pr.font.name = FONT_TECH
    pr.font.size = Pt(8.5)
    pr.font.bold = True
    pr.font.color.rgb = TEXT_DIM
    
    # Footer
    tb_f = s.shapes.add_textbox(Inches(0.8), Inches(7.12), Inches(11.733), Inches(0.25))
    tf_f = tb_f.text_frame
    tf_f.margin_left = tf_f.margin_top = tf_f.margin_right = tf_f.margin_bottom = 0
    pf = tf_f.paragraphs[0]
    pf.text = "ATMANIRBHAR DEFENCE TECH  •  HEATBOLT SOVEREIGN THERMAL PLATFORM  •  RESTRICTED BRIEFING"
    pf.font.name = FONT_TECH
    pf.font.size = Pt(7.5)
    pf.font.color.rgb = TEXT_DIM
    
    return s

def add_card(s, x, y, w, h, title, desc="", tag=None, accent_color=GOLD_SOVEREIGN, title_size=11.5, desc_size=9, bg_col=PANEL_WHITE):
    c = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    c.fill.solid()
    c.fill.fore_color.rgb = bg_col
    c.line.color.rgb = BORDER_SUBTLE
    c.line.width = Pt(1)
    
    # Top accent line
    top_line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + 0.08), Inches(y), Inches(w - 0.16), Inches(0.035))
    top_line.fill.solid()
    top_line.fill.fore_color.rgb = accent_color
    top_line.line.fill.background()
    
    tb = s.shapes.add_textbox(Inches(x + 0.16), Inches(y + 0.12), Inches(w - 0.32), Inches(h - 0.22))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    if tag:
        ptag = tf.paragraphs[0]
        ptag.text = tag.upper()
        ptag.font.name = FONT_TECH
        ptag.font.size = Pt(8)
        ptag.font.bold = True
        ptag.font.color.rgb = accent_color
        ptag.space_after = Pt(2)
        ptitle = tf.add_paragraph()
    else:
        ptitle = tf.paragraphs[0]
        
    ptitle.text = title
    ptitle.font.name = FONT_HEAD
    ptitle.font.size = Pt(title_size)
    ptitle.font.bold = True
    ptitle.font.color.rgb = NAVY_PRIMARY
    ptitle.space_after = Pt(3)
    
    if desc:
        pdesc = tf.add_paragraph()
        pdesc.text = desc
        pdesc.font.name = FONT_BODY
        pdesc.font.size = Pt(desc_size)
        pdesc.font.color.rgb = TEXT_MUTED
        pdesc.space_before = Pt(2)

def add_tag_pills(s, x_start, y, tag_list, max_w=11.7):
    cur_x = x_start
    cur_y = y
    for tag in tag_list:
        w = max(0.9, len(tag) * 0.085 + 0.28)
        if cur_x + w > x_start + max_w:
            cur_x = x_start
            cur_y += 0.36
            
        pill = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cur_x), Inches(cur_y), Inches(w), Inches(0.28))
        pill.fill.solid()
        pill.fill.fore_color.rgb = TAG_BG_LIGHT
        pill.line.color.rgb = BORDER_SUBTLE
        pill.line.width = Pt(0.75)
        
        tb = s.shapes.add_textbox(Inches(cur_x), Inches(cur_y + 0.04), Inches(w), Inches(0.2))
        tf = tb.text_frame
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = tag
        p.font.name = FONT_TECH
        p.font.size = Pt(8)
        p.font.bold = True
        p.font.color.rgb = NAVY_SECONDARY
        
        cur_x += w + 0.12
    return cur_y + 0.36

def add_note_box(s, x, y, w, h, bold_prefix, note_text, accent=GOLD_SOVEREIGN):
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = GOLD_BG_TINT
    box.line.color.rgb = GOLD_BORDER
    box.line.width = Pt(1)
    
    # Left vertical accent
    left_bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y + 0.06), Inches(0.05), Inches(h - 0.12))
    left_bar.fill.solid()
    left_bar.fill.fore_color.rgb = accent
    left_bar.line.fill.background()
    
    tb = s.shapes.add_textbox(Inches(x + 0.18), Inches(y + 0.1), Inches(w - 0.3), Inches(h - 0.2))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    
    r1 = p.add_run()
    r1.text = bold_prefix + "  "
    r1.font.name = FONT_HEAD
    r1.font.bold = True
    r1.font.size = Pt(9.5)
    r1.font.color.rgb = accent
    
    r2 = p.add_run()
    r2.text = note_text
    r2.font.name = FONT_BODY
    r2.font.size = Pt(9)
    r2.font.color.rgb = TEXT_DARK

# ==============================================================================
# SLIDE 1: COVER SLIDE (LIGHT EXECUTIVE)
# ==============================================================================
s1 = prs.slides.add_slide(blank_layout)
bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
bg1.fill.solid()
bg1.fill.fore_color.rgb = BG_CANVAS
bg1.line.fill.background()

# Frame
frame = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9))
frame.fill.solid()
frame.fill.fore_color.rgb = PANEL_WHITE
frame.line.color.rgb = BORDER_SUBTLE
frame.line.width = Pt(1.5)

# Corner Accent Brackets
c1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.8), Inches(0.4), Inches(0.06))
c1.fill.solid(); c1.fill.fore_color.rgb = GOLD_SOVEREIGN; c1.line.fill.background()
c2 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.8), Inches(0.06), Inches(0.4))
c2.fill.solid(); c2.fill.fore_color.rgb = GOLD_SOVEREIGN; c2.line.fill.background()

# Badge
badge = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.3), Inches(1.25), Inches(5.2), Inches(0.38))
badge.fill.solid()
badge.fill.fore_color.rgb = GOLD_BG_TINT
badge.line.color.rgb = GOLD_SOVEREIGN
badge.line.width = Pt(1)
tb_badge = s1.shapes.add_textbox(Inches(1.3), Inches(1.31), Inches(5.2), Inches(0.3))
tf_badge = tb_badge.text_frame
tf_badge.margin_left = tf_badge.margin_top = tf_badge.margin_right = tf_badge.margin_bottom = 0
pb = tf_badge.paragraphs[0]
pb.alignment = PP_ALIGN.CENTER
pb.text = "ATMANIRBHAR BHARAT  //  DEFENCE STRATEGIC BRIEFING"
pb.font.name = FONT_TECH
pb.font.size = Pt(9.5)
pb.font.bold = True
pb.font.color.rgb = GOLD_SOVEREIGN

# Title
tb_t1 = s1.shapes.add_textbox(Inches(1.3), Inches(1.85), Inches(10.5), Inches(1.35))
tf_t1 = tb_t1.text_frame
tf_t1.margin_left = tf_t1.margin_top = tf_t1.margin_right = tf_t1.margin_bottom = 0
pt1 = tf_t1.paragraphs[0]
pt1.text = "HEATBOLT"
pt1.font.name = FONT_HEAD
pt1.font.size = Pt(66)
pt1.font.bold = True
pt1.font.color.rgb = NAVY_PRIMARY

# Subtitle
tb_sub = s1.shapes.add_textbox(Inches(1.3), Inches(3.3), Inches(10.5), Inches(0.55))
tf_sub = tb_sub.text_frame
tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
ps = tf_sub.paragraphs[0]
ps.text = "SOVEREIGN PORTABLE THERMAL INFRASTRUCTURE FOR EXTREME ENVIRONMENTS"
ps.font.name = FONT_HEAD
ps.font.size = Pt(16)
ps.font.bold = True
ps.font.color.rgb = GOLD_SOVEREIGN

# Divider
line1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.3), Inches(3.95), Inches(3.0), Inches(0.04))
line1.fill.solid()
line1.fill.fore_color.rgb = NAVY_PRIMARY
line1.line.fill.background()

# Lede
tb_l1 = s1.shapes.add_textbox(Inches(1.3), Inches(4.15), Inches(10.5), Inches(1.1))
tf_l1 = tb_l1.text_frame
tf_l1.word_wrap = True
tf_l1.margin_left = tf_l1.margin_top = tf_l1.margin_right = tf_l1.margin_bottom = 0
pl1 = tf_l1.paragraphs[0]
pl1.text = "An indigenous, water-activated thermal platform engineered to deliver instant heat and steam without mains electricity, LPG cylinders, or open flame.\nEngineered for high-altitude border supremacy (Siachen / Eastern Ladakh / LAC), combat casualty preservation, and frontline soldier survival where conventional infrastructure cannot follow."
pl1.font.name = FONT_BODY
pl1.font.size = Pt(12)
pl1.font.color.rgb = TEXT_MUTED

# Metadata Footer
tb_m1 = s1.shapes.add_textbox(Inches(1.3), Inches(5.6), Inches(6.0), Inches(0.7))
tf_m1 = tb_m1.text_frame
tf_m1.margin_left = tf_m1.margin_top = tf_m1.margin_right = tf_m1.margin_bottom = 0
pm1 = tf_m1.paragraphs[0]
pm1.text = "DR. KUMARPAL PORWAL — FOUNDER & MANAGING DIRECTOR\n+91 73533 19757 | founder.heatbolt@gmail.com | Bangalore, India\nDefence Innovation & Technology Commercialisation"
pm1.font.name = FONT_TECH
pm1.font.size = Pt(8.5)
pm1.font.color.rgb = TEXT_MUTED

tb_m2 = s1.shapes.add_textbox(Inches(7.2), Inches(5.6), Inches(4.8), Inches(0.7))
tf_m2 = tb_m2.text_frame
tf_m2.margin_left = tf_m2.margin_top = tf_m2.margin_right = tf_m2.margin_bottom = 0
pm2 = tf_m2.paragraphs[0]
pm2.alignment = PP_ALIGN.RIGHT
pm2.text = "MINISTRY OF DEFENCE & ARMED FORCES BRIEFING\nHigh-Altitude & Winter Warfare Doctrine\n100% Make-in-India Intellectual Property"
pm2.font.name = FONT_TECH
pm2.font.size = Pt(8.5)
pm2.font.color.rgb = GOLD_SOVEREIGN

print("Slide 1 generated.")

# ==============================================================================
# SLIDE 2: THE ONE-LINE PROPOSITION
# ==============================================================================
s2 = base_slide("The One-Line Proposition", "Thermal Infrastructure at the Point of Need", "02 / 24")

# Vision Card
c_vis = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.65), Inches(11.733), Inches(1.4))
c_vis.fill.solid()
c_vis.fill.fore_color.rgb = PANEL_WHITE
c_vis.line.color.rgb = GOLD_BORDER
c_vis.line.width = Pt(1.5)

tb_v = s2.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(11.1), Inches(1.1))
tf_v = tb_v.text_frame
tf_v.word_wrap = True
tf_v.margin_left = tf_v.margin_top = tf_v.margin_right = tf_v.margin_bottom = 0
pv_k = tf_v.paragraphs[0]
pv_k.text = "// STRATEGIC DEFENCE DOCTRINE"
pv_k.font.name = FONT_TECH
pv_k.font.size = Pt(9.5)
pv_k.font.bold = True
pv_k.font.color.rgb = GOLD_SOVEREIGN
pv_k.space_after = Pt(4)

pv_m = tf_v.add_paragraph()
pv_m.text = "Move thermal infrastructure from the base to the soldier."
pv_m.font.name = FONT_HEAD
pv_m.font.size = Pt(19)
pv_m.font.bold = True
pv_m.font.color.rgb = NAVY_PRIMARY

# Core Proposition Lede
tb_l2 = s2.shapes.add_textbox(Inches(0.8), Inches(3.25), Inches(11.733), Inches(1.2))
tf_l2 = tb_l2.text_frame
tf_l2.word_wrap = True
tf_l2.margin_left = tf_l2.margin_top = tf_l2.margin_right = tf_l2.margin_bottom = 0
pl2 = tf_l2.paragraphs[0]
pl2.text = "Heatbolt is developing an indigenous, water-activated thermal platform providing controlled heat and steam without dependence on fixed infrastructure — designed for frontline soldier warmth, hot ration hydration, high-altitude alpine warfare, combat casualty support, and forward shelters."
pl2.font.name = FONT_BODY
pl2.font.size = Pt(13)
pl2.font.color.rgb = TEXT_MUTED

# Tag Pills
tag_list_s2 = [
    "Soldier Warmth", "Hot Combat Rations", "High-Altitude Ops",
    "Snow-to-Water Conversion", "Combat Casualty Care", "Field Shelters",
    "Hygiene & Sanitation", "Weapon De-Icing", "UAV Subsystems", "Search & Rescue"
]
add_tag_pills(s2, 0.8, 4.65, tag_list_s2, max_w=11.7)

# Guiding Note Box
add_note_box(s2, 0.8, 5.75, 11.733, 0.95, "SOVEREIGN DEFENCE VALUE —", "Every frontline soldier must possess dependable thermal independence. Frontline endurance, physical combat agility, cognitive acuity, and casualty survival in -40°C glacial theatres depend directly on immediate, localized thermal energy.")

print("Slide 2 generated.")

# ==============================================================================
# SLIDE 3: SECTION 01 DIVIDER
# ==============================================================================
s3 = base_slide("Section 01 // Operational Context", "Extreme Environments Create a Thermal Infrastructure Gap", "03 / 24")

c_s3 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.8))
c_s3.fill.solid()
c_s3.fill.fore_color.rgb = PANEL_WHITE
c_s3.line.color.rgb = BORDER_SUBTLE
c_s3.line.width = Pt(1)

tb_s3 = s3.shapes.add_textbox(Inches(1.3), Inches(2.3), Inches(10.7), Inches(3.8))
tf_s3 = tb_s3.text_frame
tf_s3.word_wrap = True
tf_s3.margin_left = tf_s3.margin_top = tf_s3.margin_right = tf_s3.margin_bottom = 0

ps3_1 = tf_s3.paragraphs[0]
ps3_1.text = "THE BORDER & HIGH-ALTITUDE LOGISTICS BOTTLENECK"
ps3_1.font.name = FONT_TECH
ps3_1.font.size = Pt(11)
ps3_1.font.bold = True
ps3_1.font.color.rgb = GOLD_SOVEREIGN
ps3_1.space_after = Pt(10)

ps3_2 = tf_s3.add_paragraph()
ps3_2.text = "A soldier in high-altitude, remote, glacial, winter or infrastructure-constrained terrain needs thermal energy for far more than personal warmth."
ps3_2.font.name = FONT_HEAD
ps3_2.font.size = Pt(22)
ps3_2.font.bold = True
ps3_2.font.color.rgb = NAVY_PRIMARY
ps3_2.space_after = Pt(16)

ps3_3 = tf_s3.add_paragraph()
ps3_3.text = "Modern military logistics still rely heavily on heavy diesel generators, vulnerable fuel bladders, and volatile kerosene stoves for forward mountain positions. When patrol detachments deploy to isolated ridges, observation posts, and glacial passes, external infrastructure is virtually non-existent.\n\nWithout portable, flameless, solid-state thermal capability, hypothermia rates spike, combat rations freeze into unpalatable blocks, medical triage is severely compromised, and weapon linkages seize in critical moments."
ps3_3.font.name = FONT_BODY
ps3_3.font.size = Pt(12)
ps3_3.font.color.rgb = TEXT_MUTED

print("Slide 3 generated.")

# ==============================================================================
# SLIDE 4: THE OPERATIONAL PROBLEM (8 CARDS)
# ==============================================================================
s4 = base_slide("The Operational Problem", "Eight Thermal Requirements. Zero Infrastructure.", "04 / 24")

reqs = [
    ("01 Personnel Warmth", "Sustained warmth, frostbite prevention, survival, and cold-weather physical endurance."),
    ("02 Combat Rations", "Flameless heating of RTE / RTC rations, hot tea, and vital warm hydration on patrol."),
    ("03 Water Utility", "Melting high-altitude snowpack into hot drinking water and utility steam."),
    ("04 Field Medical Care", "Casualty hypothermia prevention during evacuation and temperature-controlled medical triage."),
    ("05 Equipment Readiness", "Drying tactical gear, unfreezing weapon breach blocks, de-icing optics, and vehicle start."),
    ("06 Shelter Heating", "Temporary heating of alpine bivouac shelters, tents, and forward observation posts."),
    ("07 Tactical Mobility", "Cold-weather pre-heating for vehicle engine blocks, optics, and UAV lithium batteries."),
    ("08 Emergency SAR", "High-altitude search & rescue, avalanche survivor extraction, and disaster relief.")
]

card_w = 2.75
card_h = 2.2
for idx, (title, desc) in enumerate(reqs):
    row = idx // 4
    col = idx % 4
    cx = 0.8 + col * (card_w + 0.24)
    cy = 1.7 + row * (card_h + 0.3)
    add_card(s4, cx, cy, card_w, card_h, title, desc, tag=f"DEMAND #{idx+1:02d}", accent_color=GOLD_SOVEREIGN)

print("Slide 4 generated.")

# ==============================================================================
# SLIDE 5: THE CHALLENGE
# ==============================================================================
s5 = base_slide("The Tactical Conundrum", "The Challenge of Useful Heat", "05 / 24")

q_box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(11.733), Inches(2.2))
q_box.fill.solid()
q_box.fill.fore_color.rgb = GOLD_BG_TINT
q_box.line.color.rgb = GOLD_BORDER
q_box.line.width = Pt(1.5)

tb_q = s5.shapes.add_textbox(Inches(1.2), Inches(1.9), Inches(10.9), Inches(1.8))
tf_q = tb_q.text_frame
tf_q.word_wrap = True
tf_q.margin_left = tf_q.margin_top = tf_q.margin_right = tf_q.margin_bottom = 0

pq1 = tf_q.paragraphs[0]
pq1.text = "THE FUNDAMENTAL COMBAT PARADOX"
pq1.font.name = FONT_TECH
pq1.font.size = Pt(9.5)
pq1.font.bold = True
pq1.font.color.rgb = GOLD_SOVEREIGN
pq1.space_after = Pt(6)

pq2 = tf_q.add_paragraph()
pq2.text = "“The challenge is not simply producing heat.\nIt is producing USEFUL HEAT where conventional infrastructure cannot easily follow.”"
pq2.font.name = FONT_HEAD
pq2.font.size = Pt(20)
pq2.font.bold = True
pq2.font.color.rgb = NAVY_PRIMARY

tags_s5 = ["Sub-Zero (-40°C)", "Remote Mountain Ridge", "Blackout Operations", "Zero Logistical Trail", "Physically Demanding", "Zero Carbon Monoxide"]
add_tag_pills(s5, 0.8, 4.25, tags_s5, max_w=11.7)

c_ep = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.0), Inches(11.733), Inches(1.7))
c_ep.fill.solid()
c_ep.fill.fore_color.rgb = PANEL_WHITE
c_ep.line.color.rgb = BORDER_SUBTLE
c_ep.line.width = Pt(1)

tb_ep = s5.shapes.add_textbox(Inches(1.1), Inches(5.2), Inches(11.1), Inches(1.3))
tf_ep = tb_ep.text_frame
tf_ep.word_wrap = True
tf_ep.margin_left = tf_ep.margin_top = tf_ep.margin_right = tf_ep.margin_bottom = 0

pep1 = tf_ep.paragraphs[0]
pep1.text = "THE FRONTLINE SOLDIER AS THE BENCHMARK FOR EXTREME ENGINEERING"
pep1.font.name = FONT_TECH
pep1.font.size = Pt(9)
pep1.font.bold = True
pep1.font.color.rgb = GOLD_SOVEREIGN
pep1.space_after = Pt(4)

pep2 = tf_ep.add_paragraph()
pep2.text = "The frontline soldier operating in glacial altitude represents the ultimate operating stress test. If a thermal platform performs reliably under high-altitude, wind-chill, and sub-zero conditions, every adjacent commercial, industrial, and humanitarian application becomes a natural expansion — not an unvalidated risk."
pep2.font.name = FONT_BODY
pep2.font.size = Pt(11)
pep2.font.color.rgb = TEXT_MUTED

print("Slide 5 generated.")

# ==============================================================================
# SLIDE 6: SECTION 02 DIVIDER
# ==============================================================================
s6 = base_slide("Section 02 // Platform Architecture", "Portable Thermal Infrastructure", "06 / 24")

c_s6 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.8))
c_s6.fill.solid()
c_s6.fill.fore_color.rgb = PANEL_WHITE
c_s6.line.color.rgb = BORDER_SUBTLE
c_s6.line.width = Pt(1)

tb_s6 = s6.shapes.add_textbox(Inches(1.3), Inches(2.3), Inches(10.7), Inches(3.8))
tf_s6 = tb_s6.text_frame
tf_s6.word_wrap = True
tf_s6.margin_left = tf_s6.margin_top = tf_s6.margin_right = tf_s6.margin_bottom = 0

ps6_1 = tf_s6.paragraphs[0]
ps6_1.text = "100% INDIGENOUS REACTION ENGINEERING"
ps6_1.font.name = FONT_TECH
ps6_1.font.size = Pt(11)
ps6_1.font.bold = True
ps6_1.font.color.rgb = GOLD_SOVEREIGN
ps6_1.space_after = Pt(10)

ps6_2 = tf_s6.add_paragraph()
ps6_2.text = "A proprietary water-activated exothermic system — engineered specifically for extreme environments where infrastructure is zero."
ps6_2.font.name = FONT_HEAD
ps6_2.font.size = Pt(22)
ps6_2.font.bold = True
ps6_2.font.color.rgb = NAVY_PRIMARY
ps6_2.space_after = Pt(16)

ps6_3 = tf_s6.add_paragraph()
ps6_3.text = "Heatbolt eliminates dangerous open fires, heavy generators, and hazardous pressurized fuel canisters through a solid-state, highly stable mineral formulation.\n\nActivated by ambient water, stream runoff, or melted snowpack, Heatbolt converts chemical bond energy directly into rapid, controlled thermodynamic heat and clean steam in seconds — with 100% indigenous Indian raw materials."
ps6_3.font.name = FONT_BODY
ps6_3.font.size = Pt(12)
ps6_3.font.color.rgb = TEXT_MUTED

print("Slide 6 generated.")

# ==============================================================================
# SLIDE 7: THE HEATBOLT ANSWER (4-STEP FLOW)
# ==============================================================================
s7 = base_slide("The Heatbolt Answer", "Water In. Heat Out. Anywhere.", "07 / 24")

steps = [
    ("01", "ACTIVATE", "Water or melted snow is introduced into the hermetic thermal pouch chamber."),
    ("02", "REACT", "Proprietary mineral formulation triggers an instant, controlled exothermic reaction."),
    ("03", "TRANSFER", "Thermal energy and pure steam conduct precisely into the payload without contamination."),
    ("04", "DEPLOY", "Engineered into soldier pouches, ration packs, survival vests, and tactical equipment.")
]

sw = 2.75
sh = 3.2
for idx, (num, step_name, step_desc) in enumerate(steps):
    sx = 0.8 + idx * (sw + 0.24)
    sy = 1.75
    
    sc = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(sx), Inches(sy), Inches(sw), Inches(sh))
    sc.fill.solid()
    sc.fill.fore_color.rgb = PANEL_WHITE
    sc.line.color.rgb = BORDER_SUBTLE
    sc.line.width = Pt(1)
    
    # Top Gold line
    st_line = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(sx + 0.08), Inches(sy), Inches(sw - 0.16), Inches(0.035))
    st_line.fill.solid(); st_line.fill.fore_color.rgb = GOLD_SOVEREIGN; st_line.line.fill.background()
    
    tb_num = s7.shapes.add_textbox(Inches(sx + 0.2), Inches(sy + 0.2), Inches(sw - 0.4), Inches(0.7))
    tf_num = tb_num.text_frame
    tf_num.margin_left = tf_num.margin_top = tf_num.margin_right = tf_num.margin_bottom = 0
    pn = tf_num.paragraphs[0]
    pn.text = num
    pn.font.name = FONT_TECH
    pn.font.size = Pt(36)
    pn.font.bold = True
    pn.font.color.rgb = GOLD_SOVEREIGN
    
    tb_sd = s7.shapes.add_textbox(Inches(sx + 0.2), Inches(sy + 1.1), Inches(sw - 0.4), Inches(1.9))
    tf_sd = tb_sd.text_frame
    tf_sd.word_wrap = True
    tf_sd.margin_left = tf_sd.margin_top = tf_sd.margin_right = tf_sd.margin_bottom = 0
    psn = tf_sd.paragraphs[0]
    psn.text = step_name
    psn.font.name = FONT_HEAD
    psn.font.size = Pt(15)
    psn.font.bold = True
    psn.font.color.rgb = NAVY_PRIMARY
    psn.space_after = Pt(8)
    
    psb = tf_sd.add_paragraph()
    psb.text = step_desc
    psb.font.name = FONT_BODY
    psb.font.size = Pt(10)
    psb.font.color.rgb = TEXT_MUTED

tags_s7 = ["Zero Mains Power", "Zero LPG / Gas Need", "Zero Open Flame", "Zero Acoustic Signature", "Zero Toxic Residue", "Modular Form Factors"]
add_tag_pills(s7, 0.8, 5.35, tags_s7, max_w=11.7)

add_note_box(s7, 0.8, 6.0, 11.733, 0.8, "TACTICAL BLACKOUT DISCIPLINE —", "Flameless chemical activation emits zero light signature and zero acoustic profile, fully conforming to strict night blackout and forward sentry concealment doctrines.")

print("Slide 7 generated.")

# ==============================================================================
# SLIDE 8: THE PLATFORM THESIS
# ==============================================================================
s8 = base_slide("The Platform Thesis", "Not a Heater. A Thermal Platform.", "08 / 24")

tb_l8 = s8.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.733), Inches(0.45))
tf_l8 = tb_l8.text_frame
tf_l8.margin_left = tf_l8.margin_top = tf_l8.margin_right = tf_l8.margin_bottom = 0
pl8 = tf_l8.paragraphs[0]
pl8.text = "A conventional product solves one problem. A platform solves many — using one common technological core."
pl8.font.name = FONT_BODY
pl8.font.size = Pt(12)
pl8.font.color.rgb = GOLD_SOVEREIGN

p_cards = [
    ("Soldier Warming", "Personal thermal protection at point of need; wearable heating elements."),
    ("Ration Heating", "Hot meals and beverages without field kitchens, flame, or combustion smoke."),
    ("Hot Water / Steam", "Field utility hot water generation, snow-to-water conversion, sanitation."),
    ("Temporary Heating", "Bivouac shelters, observation posts, forward operational living pods."),
    ("Field Medical", "Thermal workflows for casualty care, hypothermia treatment, IV warming."),
    ("Equipment Drying", "Tactical gear drying, weapon de-icing, moisture mitigation in freezing zones."),
    ("Cold-Weather Mobility", "Battery pre-heating, UAV sub-systems, vehicle component readiness."),
    ("Emergency SAR", "Disaster response, avalanche search & rescue, humanitarian mission thermal kits.")
]

cw8 = 2.75
ch8 = 2.05
for idx, (t, d) in enumerate(p_cards):
    r = idx // 4
    c = idx % 4
    x = 0.8 + c * (cw8 + 0.24)
    y = 2.1 + r * (ch8 + 0.28)
    add_card(s8, x, y, cw8, ch8, t, d, tag=f"DOMAIN #{idx+1:02d}", accent_color=GOLD_SOVEREIGN)

print("Slide 8 generated.")

# ==============================================================================
# SLIDE 9: PERFORMANCE ENVELOPE (MASSIVE STATS)
# ==============================================================================
s9 = base_slide("Performance Envelope", "Selected Figures from Current Development", "09 / 24")

metrics = [
    ("100°C", "Peak Thermal Output", "Selected high-performance configurations"),
    ("< 30s", "Activation Latency", "Approximate rapid-response trigger"),
    ("45 min", "Active Thermal Duration", "Standard military operational envelope"),
    ("60°C / 3h", "Extended Warmth Profile", "Casualty transport & survival configuration")
]

mw = 2.75
mh = 3.4
for idx, (num, lbl, sub) in enumerate(metrics):
    mx = 0.8 + idx * (mw + 0.24)
    my = 1.8
    
    mc = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(mx), Inches(my), Inches(mw), Inches(mh))
    mc.fill.solid()
    mc.fill.fore_color.rgb = PANEL_WHITE
    mc.line.color.rgb = GOLD_BORDER
    mc.line.width = Pt(1.5)
    
    m_top = s9.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(mx + 0.1), Inches(my), Inches(mw - 0.2), Inches(0.04))
    m_top.fill.solid(); m_top.fill.fore_color.rgb = GOLD_SOVEREIGN; m_top.line.fill.background()
    
    tb_m = s9.shapes.add_textbox(Inches(mx + 0.15), Inches(my + 0.4), Inches(mw - 0.3), Inches(2.6))
    tf_m = tb_m.text_frame
    tf_m.word_wrap = True
    tf_m.margin_left = tf_m.margin_top = tf_m.margin_right = tf_m.margin_bottom = 0
    
    pm1 = tf_m.paragraphs[0]
    pm1.alignment = PP_ALIGN.CENTER
    pm1.text = num
    pm1.font.name = FONT_HEAD
    pm1.font.size = Pt(36)
    pm1.font.bold = True
    pm1.font.color.rgb = GOLD_SOVEREIGN
    pm1.space_after = Pt(12)
    
    pm2 = tf_m.add_paragraph()
    pm2.alignment = PP_ALIGN.CENTER
    pm2.text = lbl
    pm2.font.name = FONT_HEAD
    pm2.font.size = Pt(12)
    pm2.font.bold = True
    pm2.font.color.rgb = NAVY_PRIMARY
    pm2.space_after = Pt(8)
    
    pm3 = tf_m.add_paragraph()
    pm3.alignment = PP_ALIGN.CENTER
    pm3.text = sub
    pm3.font.name = FONT_BODY
    pm3.font.size = Pt(9.5)
    pm3.font.color.rgb = TEXT_MUTED

add_note_box(s9, 0.8, 5.6, 11.733, 1.15, "QUALIFICATION MANDATE —", "Performance figures reflect developmental laboratory envelopes. Output kinetics vary with formulation density, water volume, packaging geometry, payload thermal conductivity, ambient sub-zero temperature, and wind chill. Every Defence application undergoes rigorous mil-spec qualification.")

print("Slide 9 generated.")

# ==============================================================================
# SLIDE 10: SECTION 03 DIVIDER
# ==============================================================================
s10 = base_slide("Section 03 // Mission Capabilities", "From Product to Soldier Capability", "10 / 24")

c_s10 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.8))
c_s10.fill.solid()
c_s10.fill.fore_color.rgb = PANEL_WHITE
c_s10.line.color.rgb = BORDER_SUBTLE
c_s10.line.width = Pt(1)

tb_s10 = s10.shapes.add_textbox(Inches(1.3), Inches(2.3), Inches(10.7), Inches(3.8))
tf_s10 = tb_s10.text_frame
tf_s10.word_wrap = True
tf_s10.margin_left = tf_s10.margin_top = tf_s10.margin_right = tf_s10.margin_bottom = 0

ps10_1 = tf_s10.paragraphs[0]
ps10_1.text = "MISSION-ALIGNED DEFENCE SOLUTIONS"
ps10_1.font.name = FONT_TECH
ps10_1.font.size = Pt(11)
ps10_1.font.bold = True
ps10_1.font.color.rgb = GOLD_SOVEREIGN
ps10_1.space_after = Pt(10)

ps10_2 = tf_s10.add_paragraph()
ps10_2.text = "Ten mission-aligned use cases — each engineered toward rigorous military qualification, not mere laboratory demonstration."
ps10_2.font.name = FONT_HEAD
ps10_2.font.size = Pt(22)
ps10_2.font.bold = True
ps10_2.font.color.rgb = NAVY_PRIMARY
ps10_2.space_after = Pt(16)

ps10_3 = tf_s10.add_paragraph()
ps10_3.text = "Heatbolt translates fundamental chemical thermodynamics into practical soldier-centric kits.\n\nFrom extreme glacial combat at -40°C on the Siachen glacier to rapid forward triage casualty evacuation, Heatbolt creates modular capability that integrates directly into existing defence procurement specifications."
ps10_3.font.name = FONT_BODY
ps10_3.font.size = Pt(12)
ps10_3.font.color.rgb = TEXT_MUTED

print("Slide 10 generated.")

# ==============================================================================
# SLIDE 11: MISSION USE CASES (10 CARDS)
# ==============================================================================
s11 = base_slide("Mission Use Cases", "Ten Use Cases. One Core Platform.", "11 / 24")

cases = [
    ("UC-01 Personal Thermal", "Body warmers, survival packs, heated gloves/jackets for prolonged sub-zero sentry watch."),
    ("UC-02 Field Food & Beverages", "Flameless self-heating of combat rations (RTE/RTC meals), tea, coffee, and clean hot water."),
    ("UC-03 High-Altitude Glacial", "Warming, hot hydration, shelter heating, gear drying, and continuous snow melting."),
    ("UC-04 Ice & Snow Management", "Localised melting, clearing frozen weapon mounts, optics, and locked access-points."),
    ("UC-05 Field Medical Support", "Casualty transport hypothermia prevention, warm IV fluids, and field sterilisation."),
    ("UC-06 Temporary Shelters", "Rapid emergency heating for tactical bivouacs, alpine shelters, and forward command posts."),
    ("UC-07 Hygiene & Sanitation", "On-demand hot water, steam generation, field laundry, and decontamination stations."),
    ("UC-08 Equipment Maintenance", "Weapon bore drying, vehicle engine block thawing, and frozen linkage recovery."),
    ("UC-09 Aviation & UAV Systems", "Pre-flight surface de-icing and cold-weather UAV battery pre-heating support."),
    ("UC-10 Search & Rescue (SAR)", "High-altitude disaster recovery, avalanche rescue thermal bags, and multi-agency response.")
]

cw11 = 2.2
ch11 = 2.1
for idx, (title, desc) in enumerate(cases):
    r = idx // 5
    c = idx % 5
    x = 0.8 + c * (cw11 + 0.18)
    y = 1.8 + r * (ch11 + 0.3)
    add_card(s11, x, y, cw11, ch11, title, desc, tag=f"USE CASE {idx+1:02d}", accent_color=GOLD_SOVEREIGN, title_size=10.5, desc_size=8.5)

print("Slide 11 generated.")

# ==============================================================================
# SLIDE 12: THE SOLDIER-FIRST VISION (4 MISSION KITS)
# ==============================================================================
s12 = base_slide("The Soldier-First Vision", "A Family of Mission-Ready Kits", "12 / 24")

kits = [
    ("KIT 01", "Soldier Survival Kit", "Portable thermal activation • Emergency body warming • Food & beverage heating • Clean hot water", "Immediate survival loadout for extreme weather patrols and covert reconnaissance detachments."),
    ("KIT 02", "High-Altitude Kit", "Personal warmth modules • Snow / ice thawing packs • Weapon / gear drying • Bivouac shelter heat", "Engineered for 15,000+ ft glacial environments, permanent snowlines, and sub-zero sentry posts."),
    ("KIT 03", "Field Ration System", "Integrated heating sleeve • Hot meals in 10 mins • Hot hydration packs • Controlled steam venting", "Self-contained combat meal heating system replacing flame stoves and unpalatable frozen food."),
    ("KIT 04", "Winter Operations Kit", "Extended body warming • Heavy gear drying • Utility hot water • Mobile field clinic heating", "Comprehensive squad-level thermal kit for extreme winter warfare and prolonged field maneuvers.")
]

kw = 2.75
kh = 4.6
for idx, (code, title, specs, desc) in enumerate(kits):
    kx = 0.8 + idx * (kw + 0.24)
    ky = 1.8
    
    kc = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(kx), Inches(ky), Inches(kw), Inches(kh))
    kc.fill.solid()
    kc.fill.fore_color.rgb = PANEL_WHITE
    kc.line.color.rgb = GOLD_BORDER
    kc.line.width = Pt(1.5)
    
    k_top = s12.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(kx + 0.1), Inches(ky), Inches(kw - 0.2), Inches(0.04))
    k_top.fill.solid(); k_top.fill.fore_color.rgb = GOLD_SOVEREIGN; k_top.line.fill.background()
    
    tb_k = s12.shapes.add_textbox(Inches(kx + 0.18), Inches(ky + 0.2), Inches(kw - 0.36), Inches(4.1))
    tf_k = tb_k.text_frame
    tf_k.word_wrap = True
    tf_k.margin_left = tf_k.margin_top = tf_k.margin_right = tf_k.margin_bottom = 0
    
    pk1 = tf_k.paragraphs[0]
    pk1.text = code
    pk1.font.name = FONT_TECH
    pk1.font.size = Pt(11)
    pk1.font.bold = True
    pk1.font.color.rgb = GOLD_SOVEREIGN
    pk1.space_after = Pt(2)
    
    pk2 = tf_k.add_paragraph()
    pk2.text = title
    pk2.font.name = FONT_HEAD
    pk2.font.size = Pt(14)
    pk2.font.bold = True
    pk2.font.color.rgb = NAVY_PRIMARY
    pk2.space_after = Pt(10)
    
    pk3 = tf_k.add_paragraph()
    pk3.text = "CORE CAPABILITIES:"
    pk3.font.name = FONT_TECH
    pk3.font.size = Pt(8.5)
    pk3.font.bold = True
    pk3.font.color.rgb = GOLD_SOVEREIGN
    pk3.space_after = Pt(4)
    
    pk4 = tf_k.add_paragraph()
    pk4.text = specs
    pk4.font.name = FONT_BODY
    pk4.font.size = Pt(9.5)
    pk4.font.color.rgb = NAVY_PRIMARY
    pk4.space_after = Pt(12)
    
    pk5 = tf_k.add_paragraph()
    pk5.text = "MISSION PROFILE:"
    pk5.font.name = FONT_TECH
    pk5.font.size = Pt(8.5)
    pk5.font.bold = True
    pk5.font.color.rgb = GOLD_SOVEREIGN
    pk5.space_after = Pt(4)
    
    pk6 = tf_k.add_paragraph()
    pk6.text = desc
    pk6.font.name = FONT_BODY
    pk6.font.size = Pt(9)
    pk6.font.color.rgb = TEXT_MUTED

print("Slide 12 generated.")

# ==============================================================================
# SLIDE 13: UC-03 DEEP DIVE (HIGH-ALTITUDE & GLACIAL)
# ==============================================================================
s13 = base_slide("Use Case 03 — Deep Dive", "High-Altitude & Glacial Operations", "13 / 24")

c_left = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(6.0), Inches(4.9))
c_left.fill.solid()
c_left.fill.fore_color.rgb = PANEL_WHITE
c_left.line.color.rgb = BORDER_SUBTLE
c_left.line.width = Pt(1)

tb_l = s13.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.4), Inches(4.4))
tf_l = tb_l.text_frame
tf_l.word_wrap = True
tf_l.margin_left = tf_l.margin_top = tf_l.margin_right = tf_l.margin_bottom = 0

pl1 = tf_l.paragraphs[0]
pl1.text = "GLACIAL THEATRE CAPABILITIES"
pl1.font.name = FONT_TECH
pl1.font.size = Pt(9.5)
pl1.font.bold = True
pl1.font.color.rgb = GOLD_SOVEREIGN
pl1.space_after = Pt(10)

gl_items = [
    ("Soldier Sustained Warming", "Continuous personal thermal protection in -40°C blizzards; prevents frostbite and peripheral vasoconstriction."),
    ("Hot Food & Instant Hydration", "Rapid heating of canned and pouch rations without stoves; eliminates reliance on frozen solid rations."),
    ("Shelter & Bivouac Heating", "Safe, non-toxic flameless warmth for alpine tents without dangerous carbon monoxide buildup."),
    ("Snow & Ice Conversion", "Converts glacial snow into warm potable drinking water in minutes; prevents severe dehydration at altitude.")
]

for t, d in gl_items:
    pt = tf_l.add_paragraph()
    pt.text = f"•  {t}"
    pt.font.name = FONT_HEAD
    pt.font.size = Pt(11.5)
    pt.font.bold = True
    pt.font.color.rgb = NAVY_PRIMARY
    pt.space_before = Pt(6)
    
    pd = tf_l.add_paragraph()
    pd.text = d
    pd.font.name = FONT_BODY
    pd.font.size = Pt(9.5)
    pd.font.color.rgb = TEXT_MUTED

c_right = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.1), Inches(1.8), Inches(5.433), Inches(4.9))
c_right.fill.solid()
c_right.fill.fore_color.rgb = PANEL_WHITE
c_right.line.color.rgb = GOLD_BORDER
c_right.line.width = Pt(1.5)

tb_r = s13.shapes.add_textbox(Inches(7.4), Inches(2.0), Inches(4.8), Inches(4.4))
tf_r = tb_r.text_frame
tf_r.word_wrap = True
tf_r.margin_left = tf_r.margin_top = tf_r.margin_right = tf_r.margin_bottom = 0

pr1 = tf_r.paragraphs[0]
pr1.text = "TARGET THEATRES"
pr1.font.name = FONT_TECH
pr1.font.size = Pt(9.5)
pr1.font.bold = True
pr1.font.color.rgb = GOLD_SOVEREIGN
pr1.space_after = Pt(10)

pr2 = tf_r.add_paragraph()
pr2.text = "Siachen Glacier  •  Eastern Ladakh  •  Tawang Sector  •  Drass & Kargil Passes  •  High-Altitude Forward Posts"
pr2.font.name = FONT_HEAD
pr2.font.size = Pt(12)
pr2.font.bold = True
pr2.font.color.rgb = NAVY_PRIMARY
pr2.space_after = Pt(18)

pr3 = tf_r.add_paragraph()
pr3.text = "WHY HIGH-ALTITUDE LEADS THE STRATEGY:"
pr3.font.name = FONT_TECH
pr3.font.size = Pt(9)
pr3.font.bold = True
pr3.font.color.rgb = GOLD_SOVEREIGN
pr3.space_after = Pt(6)

pr4 = tf_r.add_paragraph()
pr4.text = "“High-altitude operations represent the harshest thermal infrastructure gap in modern Defence.\n\nIf Heatbolt performs at 18,000 feet under -40°C extreme winds, every adjacent ground environment becomes an operational expansion — not an unproven risk.”"
pr4.font.name = FONT_HEAD
pr4.font.size = Pt(13)
pr4.font.bold = True
pr4.font.color.rgb = NAVY_PRIMARY

print("Slide 13 generated.")

# ==============================================================================
# SLIDE 14: UC-05 DEEP DIVE (FIELD MEDICAL SUPPORT)
# ==============================================================================
s14 = base_slide("Use Case 05 — Deep Dive", "Combat Medical Support — Qualification Discipline", "14 / 24")

med_cards = [
    ("Patient Evacuation Care", "Continuous thermal support during casualty evacuation (CASEVAC); prevents hypothermic coagulopathy in wounded personnel.", "TRIAGE CRITICAL"),
    ("Hot Water & Steam Utility", "Immediate hot water for trauma cleaning, patient re-warming baths, and medical kit reconstitution.", "UTILITY SUPPORT"),
    ("Sterilisation Concepts", "Flameless high-temperature steam generation for field instrument sterilisation and emergency decontamination workflows.", "FIELD STERILISATION")
]

mw14 = 3.75
mh14 = 3.2
for idx, (t, d, tag) in enumerate(med_cards):
    mx = 0.8 + idx * (mw14 + 0.24)
    my = 1.8
    add_card(s14, mx, my, mw14, mh14, t, d, tag=tag, accent_color=CRIMSON_RED, title_size=13, desc_size=10)

add_note_box(s14, 0.8, 5.3, 11.733, 1.4, "CLINICAL REGULATORY DISCIPLINE —", "Heatbolt does not propose bypassing medical qualification protocols. Dedicated clinical evaluation programs cover precise thermal stability, biocompatibility, non-toxic emissions, zero volatile organic compounds, human factors, and regulatory approvals. Medical applications represent specialized development tracks.", accent=CRIMSON_RED)

print("Slide 14 generated.")

# ==============================================================================
# SLIDE 15: UC-09 & 10 (AVIATION, UAV & SAR DUAL-USE)
# ==============================================================================
s15 = base_slide("Use Cases 09 / 10 — Development Horizons", "Aviation, UAV & Emergency Response", "15 / 24")

add_card(s15, 0.8, 1.8, 5.74, 4.8, "Aviation & UAV — Future Concepts", 
         "Exploratory applications focusing on cold-weather aviation readiness:\n\n"
         "• Aircraft De-Icing Support: Rapid thermal surface thawing for rotary blades and wings.\n"
         "• UAV Battery Pre-Heating: Lithium-ion battery conditioning to restore operational flight time in sub-zero air.\n"
         "• Subsystem Warming: Thermal management for high-altitude sensor gimbals and forward-looking infrared optics.\n"
         "• Field Maintenance: Rapid thawing of seized linkages and hydraulic junctions.",
         tag="AEROSPACE HORIZONS", accent_color=NAVY_PRIMARY, title_size=14, desc_size=10.5)

add_card(s15, 6.78, 1.8, 5.75, 4.8, "Search & Rescue — Dual-Use Value",
         "One unified platform serving Defence, Disaster Response, and Humanitarian Relief:\n\n"
         "• Rapid Casualty Warming: Rapid thermal blankets for hypothermic survivors.\n"
         "• Emergency Shelters: Instant heat for avalanche rescue camps and disaster survivor zones.\n"
         "• Dual-Use National Framework: Deployment across Indian Army, NDRF, SDRF, and global humanitarian agencies.\n"
         "• Sovereign Resilience: Stockpiled emergency thermal reserves requiring zero maintenance or batteries.",
         tag="DUAL-USE HUMANITARIAN", accent_color=GOLD_SOVEREIGN, title_size=14, desc_size=10.5)

print("Slide 15 generated.")

# ==============================================================================
# SLIDE 16: INNOVATION PIPELINE (80+ CONCEPTS / 9 DOMAINS)
# ==============================================================================
s16 = base_slide("Innovation Pipeline", "One Technology. 80+ Application Concepts.", "16 / 24")

tb_l16 = s16.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.733), Inches(0.35))
tf_l16 = tb_l16.text_frame
tf_l16.margin_left = tf_l16.margin_top = tf_l16.margin_right = tf_l16.margin_bottom = 0
pl16 = tf_l16.paragraphs[0]
pl16.text = "A structured innovation pipeline across 9 operational domains — representing structured R&D horizons, not premature claims."
pl16.font.name = FONT_BODY
pl16.font.size = Pt(11)
pl16.font.color.rgb = GOLD_SOVEREIGN

domains = [
    ("Personnel", "Body warmers • Warmer jackets • Survival packs • Portable baths • Field laundry"),
    ("Food & Nutrition", "Field kitchens • Food warming • Preservation • Field bakery • Self-heating rations"),
    ("Water & Steam", "Water heating • Steam generation • Distillation • Water purification concepts"),
    ("Medical & Triage", "Field hospitals • Casualty warming • Instrument sterilisation • Trauma kits"),
    ("Equipment & Ops", "Gear drying • Vehicle maintenance • Weapon bore cleaning • De-icing systems"),
    ("Infrastructure", "Concrete curing in cold • Remote building • Emergency shelters • Underground posts"),
    ("Security & Border", "Border post warming • Surveillance camera de-fogging • Tactical fogging concepts"),
    ("Search & Rescue", "Disaster response • Avalanche survivor kits • Humanitarian emergency supply"),
    ("Environmental", "Environmental decontamination • Waste treatment thermal support • Remote field ops")
]

cw16 = 3.75
ch16 = 1.45
for idx, (t, d) in enumerate(domains):
    r = idx // 3
    c = idx % 3
    x = 0.8 + c * (cw16 + 0.24)
    y = 1.95 + r * (ch16 + 0.22)
    add_card(s16, x, y, cw16, ch16, t, d, tag=f"DOMAIN {idx+1:02d}", accent_color=GOLD_SOVEREIGN, title_size=11, desc_size=8.5)

print("Slide 16 generated.")

# ==============================================================================
# SLIDE 17: SECTION 04 DIVIDER
# ==============================================================================
s17 = base_slide("Section 04 // Leadership & Execution", "Technology Alone Does Not Create a Defence Capability", "17 / 24")

c_s17 = s17.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.8))
c_s17.fill.solid()
c_s17.fill.fore_color.rgb = PANEL_WHITE
c_s17.line.color.rgb = BORDER_SUBTLE
c_s17.line.width = Pt(1)

tb_s17 = s17.shapes.add_textbox(Inches(1.3), Inches(2.3), Inches(10.7), Inches(3.8))
tf_s17 = tb_s17.text_frame
tf_s17.word_wrap = True
tf_s17.margin_left = tf_s17.margin_top = tf_s17.margin_right = tf_s17.margin_bottom = 0

ps17_1 = tf_s17.paragraphs[0]
ps17_1.text = "EXECUTION & SCALE"
ps17_1.font.name = FONT_TECH
ps17_1.font.size = Pt(11)
ps17_1.font.bold = True
ps17_1.font.color.rgb = GOLD_SOVEREIGN
ps17_1.space_after = Pt(10)

ps17_2 = tf_s17.add_paragraph()
ps17_2.text = "Heatbolt is assembling sovereign capability across every dimension — from atomic reaction kinetics to millions of consistent field units."
ps17_2.font.name = FONT_HEAD
ps17_2.font.size = Pt(22)
ps17_2.font.bold = True
ps17_2.font.color.rgb = NAVY_PRIMARY
ps17_2.space_after = Pt(16)

ps17_3 = tf_s17.add_paragraph()
ps17_3.text = "A breakthrough technology is only as viable as its manufacturing repeatability, supply chain resilience, and institutional leadership.\n\nHeatbolt brings together seasoned industrial titans, aerospace propulsion veterans, FMCG packaging leaders, and defence strategists to bridge the gap between invention and multi-million unit military deployment."
ps17_3.font.name = FONT_BODY
ps17_3.font.size = Pt(12)
ps17_3.font.color.rgb = TEXT_MUTED

print("Slide 17 generated.")

# ==============================================================================
# SLIDE 18: CORE LEADERSHIP (4 LEADERS)
# ==============================================================================
s18 = base_slide("Executive Leadership", "Built to Execute", "18 / 24")

leaders = [
    ("Dr. Kumarpal Porwal", "FOUNDER & MANAGING DIRECTOR", "Deep-tech commercialisation • Controlled thermal activation • Enterprise AI & operational transformation (Chief of Staff, Vega Aviation) • Defence & FMCG strategic engagement.", GOLD_SOVEREIGN),
    ("Rahul Awasthi", "MANUFACTURING & QUALITY HEAD", "National Mfg & Quality Head, SBC Tanzania (Pepsi) • Ex-SLMG Coca-Cola, Amazon (led 6,000+ operations workforce), Glenmark, SABMiller, Coca-Cola.", NAVY_PRIMARY),
    ("Puneet Solanki", "MARKETING & GROWTH", "Market positioning & institutional engagement • Application & partnership development • Go-to-market scaling across Defence, institutional buyers, and global consumer brands.", GOLD_SOVEREIGN),
    ("Amit Jain", "GLOBAL BUSINESS LEADERSHIP", "Former Managing Director — L'Oréal India (First Indian MD) • Executive leadership across AkzoNobel, Coca-Cola, Viacom/MTV, ICI • Consumer innovation & global scaling.", NAVY_PRIMARY)
]

lw = 2.75
lh = 4.8
for idx, (name, role, bio, col) in enumerate(leaders):
    lx = 0.8 + idx * (lw + 0.24)
    ly = 1.8
    add_card(s18, lx, ly, lw, lh, name, bio, tag=role, accent_color=col, title_size=13, desc_size=9.5)

print("Slide 18 generated.")

# ==============================================================================
# SLIDE 19: SCIENCE & ENGINEERING ECOSYSTEM
# ==============================================================================
s19 = base_slide("Science & Engineering Ecosystem", "A Multidisciplinary Research Bench", "19 / 24")

eco_rows = [
    ("Dr. Saibalendu Sarkar — TIFR", "NANOMATERIALS & CATALYSIS", "Nanomaterials, reaction kinetics, CO2 conversion, advanced mineral characterisation and thermodynamic profiling."),
    ("Anil Kumar Tripathi", "PACKAGING & BARRIER ENGINEERING", "Packaging technology, high-barrier hermetic foils, QA/QC standards, shelf-life integrity, and new product development."),
    ("Mahantesh G. Vantamuri", "MECHANICAL & TOOLING DESIGN", "Mechanical tooling, product industrial design, CAD/CAM simulation, injection moulding, and high-speed automated assembly."),
    ("T. Koshi • Prerna Bhanot • Michele Barbazza", "STRATEGY, TALENT & EXPANSION", "Corporate governance, institutional talent scaling, international business development, and sovereign trade frameworks."),
    ("Extended Institutional Network", "DEFENCE & FOOD TECH ECOSYSTEM", "Ex-CFTRI scientists, All-India food industry leadership, Vega Helmets industrial network, Microlabs, Nilons, and advanced research labs.")
]

ew = 11.733
eh = 0.85
for idx, (name, tag, desc) in enumerate(eco_rows):
    ey = 1.8 + idx * (eh + 0.16)
    
    row_box = s19.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(ey), Inches(ew), Inches(eh))
    row_box.fill.solid()
    row_box.fill.fore_color.rgb = PANEL_WHITE
    row_box.line.color.rgb = BORDER_SUBTLE
    row_box.line.width = Pt(1)
    
    acc = s19.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(ey + 0.08), Inches(0.04), Inches(eh - 0.16))
    acc.fill.solid(); acc.fill.fore_color.rgb = GOLD_SOVEREIGN; acc.line.fill.background()
    
    tb_row = s19.shapes.add_textbox(Inches(1.0), Inches(ey + 0.12), Inches(ew - 0.3), Inches(eh - 0.24))
    tf_row = tb_row.text_frame
    tf_row.word_wrap = True
    tf_row.margin_left = tf_row.margin_top = tf_row.margin_right = tf_row.margin_bottom = 0
    
    pr1 = tf_row.paragraphs[0]
    r_name = pr1.add_run()
    r_name.text = name + "  "
    r_name.font.name = FONT_HEAD
    r_name.font.size = Pt(11)
    r_name.font.bold = True
    r_name.font.color.rgb = NAVY_PRIMARY
    
    r_tag = pr1.add_run()
    r_tag.text = f"[{tag}]  "
    r_tag.font.name = FONT_TECH
    r_tag.font.size = Pt(8.5)
    r_tag.font.bold = True
    r_tag.font.color.rgb = GOLD_SOVEREIGN
    
    pr2 = tf_row.add_paragraph()
    pr2.text = desc
    pr2.font.name = FONT_BODY
    pr2.font.size = Pt(9)
    pr2.font.color.rgb = TEXT_MUTED

print("Slide 19 generated.")

# ==============================================================================
# SLIDE 20: DISTINGUISHED ADVISORS & ECOSYSTEM
# ==============================================================================
s20 = base_slide("Distinguished Advisors & Ecosystem", "Names That Strengthen the Mission", "20 / 24")

advisors = [
    ("Nambi Narayan", "ROCKET SCIENCE & PROPULSION", "Legendary ISRO Rocket Scientist — Space propulsion pioneer, national strategic aerospace leadership.", GOLD_SOVEREIGN),
    ("Pradeep Kumar Pandey", "PACKAGING DEVELOPMENT & FMCG SCALE", "Senior Advisor — President, Packaging Development, Emami Ltd • 29+ years, Gold Medalist • Ex-Procter & Gamble, Amway, Dabur.", GOLD_SOVEREIGN),
    ("Andrew Dobringer", "GLOBAL THERMOFORMING & PACKAGING", "Senior Advisor — International Thermoforming, Precision Tooling & Advanced Packaging Systems Engineering Specialist.", GOLD_SOVEREIGN),
    ("Prof. Somnath Ganguly", "RESEARCH LEADERSHIP & MATERIALS", "Senior Advisor — 40+ years of distinguished leadership in chemical formulation, materials research, and laboratory scaling.", NAVY_PRIMARY),
    ("Ravi Tejwani", "MACHINERY & AUTOMATION", "Industrial machinery expert — Manufacturing automation, precision production equipment, and high-speed packaging lines.", NAVY_PRIMARY),
    ("R. Madhavan", "NARRATIVE & STRATEGIC ENGAGEMENT", "Acclaimed film actor & director — National communication, strategic storytelling, and high-impact institutional narrative.", NAVY_PRIMARY),
    ("Marichi Mathur", "STRATEGIC ENTERPRISE NETWORK", "Senior Advisor — Strategic growth, enterprise leadership, and national defence institutional network expansion.", NAVY_PRIMARY)
]

aw = 2.75
ah = 2.25
for idx in range(4):
    name, role, desc, col = advisors[idx]
    ax = 0.8 + idx * (aw + 0.24)
    ay = 1.8
    add_card(s20, ax, ay, aw, ah, name, desc, tag=role, accent_color=col, title_size=11, desc_size=8.5)

for idx in range(3):
    name, role, desc, col = advisors[4 + idx]
    ax = 2.29 + idx * (aw + 0.24)
    ay = 4.25
    add_card(s20, ax, ay, aw, ah, name, desc, tag=role, accent_color=col, title_size=11, desc_size=8.5)

add_note_box(s20, 0.8, 6.6, 11.733, 0.45, "GOVERNANCE NOTE —", "Advisory engagements are presented according to formal domain areas and institutional supporting documentation.")

print("Slide 20 generated.")

# ==============================================================================
# SLIDE 21 (NEW SLIDE): NATIONAL VALIDATION & INDUSTRY PROOF (IMAGE SLIDE)
# ==============================================================================
s21 = base_slide("National Credibility & Elite Endorsements", "National Titans, AMUL & Global Aviation Validation", "21 / 24")

# Insert the Image Proof Collage on Left Side
img_x = 0.8
img_y = 1.65
img_w = 6.2
img_h = 4.9

# Frame around image
img_frame = s21.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(img_x - 0.05), Inches(img_y - 0.05), Inches(img_w + 0.1), Inches(img_h + 0.1))
img_frame.fill.solid()
img_frame.fill.fore_color.rgb = PANEL_WHITE
img_frame.line.color.rgb = GOLD_BORDER
img_frame.line.width = Pt(1.5)

if os.path.exists(IMAGE_PATH):
    pic = s21.shapes.add_picture(IMAGE_PATH, Inches(img_x), Inches(img_y + 0.4), width=Inches(img_w), height=Inches(img_h - 0.5))
    
# Image Caption Bar Top
tb_ic = s21.shapes.add_textbox(Inches(img_x), Inches(img_y), Inches(img_w), Inches(0.35))
tf_ic = tb_ic.text_frame
tf_ic.margin_left = tf_ic.margin_top = tf_ic.margin_right = tf_ic.margin_bottom = 0
pic_p = tf_ic.paragraphs[0]
pic_p.alignment = PP_ALIGN.CENTER
pic_p.text = "EVIDENCE FILE: SHRI RATAN TATA CONSULTATION  •  AMUL LETTER  •  EMIRATES CATERING"
pic_p.font.name = FONT_TECH
pic_p.font.size = Pt(8.5)
pic_p.font.bold = True
pic_p.font.color.rgb = GOLD_SOVEREIGN

# Right Column: High-Impact Institutional Proof Cards
rw = 5.25
rx = 7.28

proof_items = [
    ("Shri Ratan N. Tata Direct Consultation",
     "Personal strategic consultation with India's legendary industrial patriarch, reviewing Heatbolt's self-heating technology for sovereign defence resilience and national grassroots impact.",
     "PATRIARCHAL ENDORSEMENT", GOLD_SOVEREIGN),
     
    ("AMUL (GCMMF Ltd) Official Collaboration Request",
     "Official written communication from the Purchase Division: 'We know the potential of the concept, which can bring revolution in FMCG sector, we are interested for collaboration...' — India's largest food cooperative.",
     "NATIONAL COOPERATIVE GIANT", ARMY_OLIVE),
     
    ("Emirates Flight Catering Onboarding",
     "Formal international airline vendor qualification and commercial onboarding for in-flight self-heating meals and expeditionary food solutions.",
     "GLOBAL AVIATION SUPPLY", NAVY_PRIMARY),
     
    ("Surging Cross-Sector National Inbound Demand",
     "Active high-level engagements with Bikaji Foods CFO, Tier-1 defence suppliers, and major packaging summits (Techceleration 2023 / Startup India) projecting 20B+ annual unit volume.",
     "SURGING MARKET PULL", GOLD_SOVEREIGN)
]

for idx, (p_title, p_desc, p_tag, p_col) in enumerate(proof_items):
    py = 1.65 + idx * (1.18 + 0.08)
    add_card(s21, rx, py, rw, 1.18, p_title, p_desc, tag=p_tag, accent_color=p_col, title_size=11, desc_size=8.5)

add_note_box(s21, 0.8, 6.62, 11.733, 0.45, "VERIFIED INDUSTRY ADOPTION —", "Documented traction proves Heatbolt has crossed the credibility threshold: backed by India's foremost industrial leaders, national FMCG cooperatives, and international aviation supply chains.")

print("Slide 21 (New Image Proof Slide) generated.")

# ==============================================================================
# SLIDE 22: ENDORSEMENTS & MARKET PULL
# ==============================================================================
s22 = base_slide("Endorsements & Market Pull", "The Ecosystem Is Already Responding", "22 / 24")

endorsements = [
    ("Institutional Defence Pull", "Active engagement across Indian Armed Forces command hubs, DRDO ecosystem partners, and emergency response teams for high-altitude cold trials.", "DEFENCE VALIDATION"),
    ("Global FMCG Traction", "Direct commercial, self-heating ration, and meal packaging evaluations initiated with Nestlé, Barilla, Mondelēz, PepsiCo, Unilever, and McDonald's.", "FMCG ENGAGEMENT"),
    ("Commercial Momentum", "Showcased via Startup India and Invest India platforms; over 60,000+ stakeholder touchpoints and 20B+ annual unit volume potential projected.", "SCALE TRACTION")
]

ew21 = 3.75
eh21 = 3.3
for idx, (title, desc, tag) in enumerate(endorsements):
    ex = 0.8 + idx * (ew21 + 0.24)
    ey = 1.8
    add_card(s22, ex, ey, ew21, eh21, title, desc, tag=tag, accent_color=GOLD_SOVEREIGN, title_size=13, desc_size=10)

add_note_box(s22, 0.8, 5.4, 11.733, 1.25, "PROVEN DEMAND INDICATORS —", "Heatbolt is transitioning rapidly from breakthrough laboratory formulation to mission-critical defence field evaluation and high-volume industrial packaging. Tier-1 multinational food conglomerates and defence forces have validated the immense market gap for portable, flameless, water-activated thermal energy.")

print("Slide 22 generated.")

# ==============================================================================
# SLIDE 23: STRATEGIC ROADMAP (4 DEPLOYMENT PHASES)
# ==============================================================================
s23 = base_slide("Strategic Roadmap", "Path to Operational Deployment", "23 / 24")

phases = [
    ("PHASE 01", "Field Trials", "Rigorous sub-zero & high-altitude environmental stress testing with designated frontline infantry units.", "Sub-Zero Stress Testing"),
    ("PHASE 02", "Mil-Spec Certification", "Thermal output standardization, shelf-life validation (3+ yrs), and hazardous transport compliance.", "Standards Qualification"),
    ("PHASE 03", "Volume Production", "Automated pouch filling and QC scaling with Tier-1 manufacturing partners for 10M+ units monthly.", "Automated Packaging"),
    ("PHASE 04", "Force Integration", "Standard issue inclusion into ration kits, extreme climate survival packs, and mobile field clinics.", "Standard Issue Deployment")
]

pw = 2.75
ph = 4.8
for idx, (p_code, p_title, p_desc, p_focus) in enumerate(phases):
    px = 0.8 + idx * (pw + 0.24)
    py = 1.8
    
    pc = s23.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(px), Inches(py), Inches(pw), Inches(ph))
    pc.fill.solid()
    pc.fill.fore_color.rgb = PANEL_WHITE
    pc.line.color.rgb = GOLD_BORDER
    pc.line.width = Pt(1.5)
    
    p_top = s23.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(px + 0.1), Inches(py), Inches(pw - 0.2), Inches(0.04))
    p_top.fill.solid(); p_top.fill.fore_color.rgb = GOLD_SOVEREIGN; p_top.line.fill.background()
    
    tb_p = s23.shapes.add_textbox(Inches(px + 0.18), Inches(py + 0.2), Inches(pw - 0.36), Inches(4.3))
    tf_p = tb_p.text_frame
    tf_p.word_wrap = True
    tf_p.margin_left = tf_p.margin_top = tf_p.margin_right = tf_p.margin_bottom = 0
    
    pp1 = tf_p.paragraphs[0]
    pp1.text = p_code
    pp1.font.name = FONT_TECH
    pp1.font.size = Pt(11)
    pp1.font.bold = True
    pp1.font.color.rgb = GOLD_SOVEREIGN
    pp1.space_after = Pt(4)
    
    pp2 = tf_p.add_paragraph()
    pp2.text = p_title
    pp2.font.name = FONT_HEAD
    pp2.font.size = Pt(15)
    pp2.font.bold = True
    pp2.font.color.rgb = NAVY_PRIMARY
    pp2.space_after = Pt(12)
    
    pp3 = tf_p.add_paragraph()
    pp3.text = "OPERATIONAL TARGET:"
    pp3.font.name = FONT_TECH
    pp3.font.size = Pt(8.5)
    pp3.font.bold = True
    pp3.font.color.rgb = GOLD_SOVEREIGN
    pp3.space_after = Pt(4)
    
    pp4 = tf_p.add_paragraph()
    pp4.text = p_focus
    pp4.font.name = FONT_HEAD
    pp4.font.size = Pt(11)
    pp4.font.bold = True
    pp4.font.color.rgb = NAVY_PRIMARY
    pp4.space_after = Pt(12)
    
    pp5 = tf_p.add_paragraph()
    pp5.text = "MILESTONES:"
    pp5.font.name = FONT_TECH
    pp5.font.size = Pt(8.5)
    pp5.font.bold = True
    pp5.font.color.rgb = GOLD_SOVEREIGN
    pp5.space_after = Pt(4)
    
    pp6 = tf_p.add_paragraph()
    pp6.text = p_desc
    pp6.font.name = FONT_BODY
    pp6.font.size = Pt(9.5)
    pp6.font.color.rgb = TEXT_MUTED

print("Slide 23 generated.")

# ==============================================================================
# SLIDE 24: CLOSING SLIDE (LIGHT EXECUTIVE)
# ==============================================================================
s24 = prs.slides.add_slide(blank_layout)
bg24 = s24.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
bg24.fill.solid()
bg24.fill.fore_color.rgb = BG_CANVAS
bg24.line.fill.background()

frame24 = s24.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9))
frame24.fill.solid()
frame24.fill.fore_color.rgb = PANEL_WHITE
frame24.line.color.rgb = BORDER_SUBTLE
frame24.line.width = Pt(1.5)

# Corner Accent Brackets
c24_1 = s24.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.8), Inches(0.4), Inches(0.06))
c24_1.fill.solid(); c24_1.fill.fore_color.rgb = GOLD_SOVEREIGN; c24_1.line.fill.background()
c24_2 = s24.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.8), Inches(0.06), Inches(0.4))
c24_2.fill.solid(); c24_2.fill.fore_color.rgb = GOLD_SOVEREIGN; c24_2.line.fill.background()

# Badge
badge24 = s24.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.3), Inches(1.2), Inches(5.2), Inches(0.38))
badge24.fill.solid()
badge24.fill.fore_color.rgb = GOLD_BG_TINT
badge24.line.color.rgb = GOLD_SOVEREIGN
badge24.line.width = Pt(1)
tb_b24 = s24.shapes.add_textbox(Inches(1.3), Inches(1.26), Inches(5.2), Inches(0.3))
tf_b24 = tb_b24.text_frame
tf_b24.margin_left = tf_b24.margin_top = tf_b24.margin_right = tf_b24.margin_bottom = 0
pb24 = tf_b24.paragraphs[0]
pb24.alignment = PP_ALIGN.CENTER
pb24.text = "JOIN THE MISSION  //  SOVEREIGN DEFENCE ENGAGEMENT"
pb24.font.name = FONT_TECH
pb24.font.size = Pt(9.5)
pb24.font.bold = True
pb24.font.color.rgb = GOLD_SOVEREIGN

# Title
tb_t24 = s24.shapes.add_textbox(Inches(1.3), Inches(1.85), Inches(10.5), Inches(1.2))
tf_t24 = tb_t24.text_frame
tf_t24.margin_left = tf_t24.margin_top = tf_t24.margin_right = tf_t24.margin_bottom = 0
pt24 = tf_t24.paragraphs[0]
pt24.text = "HEATBOLT"
pt24.font.name = FONT_HEAD
pt24.font.size = Pt(56)
pt24.font.bold = True
pt24.font.color.rgb = NAVY_PRIMARY

# Subtitle
tb_sub24 = s24.shapes.add_textbox(Inches(1.3), Inches(3.05), Inches(10.5), Inches(0.5))
tf_sub24 = tb_sub24.text_frame
tf_sub24.margin_left = tf_sub24.margin_top = tf_sub24.margin_right = tf_sub24.margin_bottom = 0
ps24 = tf_sub24.paragraphs[0]
ps24.text = "Partnering to Equip the Frontline Soldier with Sovereign Thermal Independence"
ps24.font.name = FONT_HEAD
ps24.font.size = Pt(17)
ps24.font.bold = True
ps24.font.color.rgb = GOLD_SOVEREIGN

# Divider
line24 = s24.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.3), Inches(3.65), Inches(3.0), Inches(0.04))
line24.fill.solid()
line24.fill.fore_color.rgb = NAVY_PRIMARY
line24.line.fill.background()

# Massive Mission Quote
tb_q24 = s24.shapes.add_textbox(Inches(1.3), Inches(3.85), Inches(10.5), Inches(1.2))
tf_q24 = tb_q24.text_frame
tf_q24.word_wrap = True
tf_q24.margin_left = tf_q24.margin_top = tf_q24.margin_right = tf_q24.margin_bottom = 0
pq24 = tf_q24.paragraphs[0]
pq24.text = "“The future of frontline combat endurance and extreme survival is not just portable food and water — it is sovereign portable heat.”"
pq24.font.name = FONT_HEAD
pq24.font.size = Pt(19)
pq24.font.bold = True
pq24.font.color.rgb = NAVY_PRIMARY

# Contact Details Left
tb_c24 = s24.shapes.add_textbox(Inches(1.3), Inches(5.3), Inches(6.0), Inches(1.1))
tf_c24 = tb_c24.text_frame
tf_c24.margin_left = tf_c24.margin_top = tf_c24.margin_right = tf_c24.margin_bottom = 0
pc24_1 = tf_c24.paragraphs[0]
pc24_1.text = "DR. KUMARPAL PORWAL — FOUNDER & MD"
pc24_1.font.name = FONT_HEAD
pc24_1.font.size = Pt(11)
pc24_1.font.bold = True
pc24_1.font.color.rgb = NAVY_PRIMARY
pc24_1.space_after = Pt(2)

pc24_2 = tf_c24.add_paragraph()
pc24_2.text = "+91 73533 19757  |  +91 88672 94189\nfounder.heatbolt@gmail.com  |  Bangalore, India"
pc24_2.font.name = FONT_TECH
pc24_2.font.size = Pt(9.5)
pc24_2.font.color.rgb = TEXT_MUTED

# Engagement Modes Right
tb_em24 = s24.shapes.add_textbox(Inches(7.2), Inches(5.3), Inches(4.8), Inches(1.1))
tf_em24 = tb_em24.text_frame
tf_em24.margin_left = tf_em24.margin_top = tf_em24.margin_right = tf_em24.margin_bottom = 0
pem1 = tf_em24.paragraphs[0]
pem1.alignment = PP_ALIGN.RIGHT
pem1.text = "DEFENCE & STRATEGIC ENGAGEMENT"
pem1.font.name = FONT_TECH
pem1.font.size = Pt(9.5)
pem1.font.bold = True
pem1.font.color.rgb = GOLD_SOVEREIGN
pem1.space_after = Pt(4)

pem2 = tf_em24.add_paragraph()
pem2.alignment = PP_ALIGN.RIGHT
pem2.text = "MIL-SPEC TRIALS (NORTHERN COMMAND / DRDO / iDEX)\nTECHNOLOGY LICENSING & IP DEFENCE INTEGRATION\nNATIONAL STRATEGIC INDUSTRIAL ALLIANCES"
pem2.font.name = FONT_TECH
pem2.font.size = Pt(9)
pem2.font.color.rgb = NAVY_SECONDARY

print("Slide 24 generated.")

# ==============================================================================
# SAVE TO BOTH DESKTOP AND REPO
# ==============================================================================
desktop_dir = os.path.expanduser("~/Desktop")
desktop_pptx = os.path.join(desktop_dir, "HEATBOLT_Defence_Strategic_Briefing.pptx")
repo_pptx = "C:/Users/My/.gemini/antigravity/scratch/career-ops/HEATBOLT_Defence_Strategic_Briefing.pptx"

prs.save(desktop_pptx)
print(f"SUCCESS: Saved World-Class Sovereign Light Presentation to Desktop at: {desktop_pptx}")

prs.save(repo_pptx)
print(f"SUCCESS: Saved World-Class Sovereign Light Presentation to Career-Ops at: {repo_pptx}")

print("24-Slide Light Theme Presentation Generated Successfully!")
