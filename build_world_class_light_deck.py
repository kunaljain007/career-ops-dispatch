import os
import sys
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

sys.stdout.reconfigure(encoding='utf-8')

print("Generating World-Class Sovereign Light Theme Presentation (.pptx)...")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# ==============================================================================
# WORLD-CLASS SOVEREIGN DEFENCE LIGHT COLOR PALETTE
# ==============================================================================
BG_CANVAS      = RGBColor(246, 248, 251)   # #F6F8FB Pristine architectural light canvas
PANEL_WHITE    = RGBColor(255, 255, 255)   # #FFFFFF Pure crisp card surface
BORDER_SUBTLE  = RGBColor(218, 226, 237)   # #DAE2ED Structural slate border
BORDER_NAVY    = RGBColor(15, 30, 54)      # #0F1E36 Deep Imperial Navy border
NAVY_PRIMARY   = RGBColor(12, 25, 46)      # #0C192E Deep Sovereign Navy (Headlines)
NAVY_SECONDARY = RGBColor(26, 44, 76)      # #1A2C4C Medium Imperial Navy
GOLD_SOVEREIGN = RGBColor(180, 130, 25)    # #B48219 Rich Atmanirbhar Gold
GOLD_BG_TINT   = RGBColor(254, 250, 240)   # #FEFAF0 Soft Gold Callout background
GOLD_BORDER    = RGBColor(220, 175, 75)    # #DCAF4B Gold Accent Border
TEXT_DARK      = RGBColor(15, 23, 42)      # #0F172A Primary text
TEXT_MUTED     = RGBColor(71, 85, 105)     # #475569 Body text
TEXT_DIM       = RGBColor(140, 155, 175)   # #8C9BAF Technical metadata
TAG_BG_LIGHT   = RGBColor(238, 243, 250)   # #EEF3FA Tag badge background
ARMY_OLIVE     = RGBColor(45, 74, 34)      # #2D4A22 Tactical Military Green
CRIMSON_RED    = RGBColor(185, 28, 28)     # #B91C1C High-priority alert red

FONT_HEAD = "Segoe UI"
FONT_BODY = "Segoe UI"
FONT_TECH = "Consolas"

blank_layout = prs.slide_layouts[6]
IMAGE_PATH = r"C:\Users\My\.gemini\antigravity\scratch\career-ops\heatbolt_national_traction.png"

def base_slide(kicker, title, slide_idx=""):
    s = prs.slides.add_slide(blank_layout)
    
    # Background Canvas
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_CANVAS
    bg.line.fill.background()
    
    # Top Sovereign Tri-Color Inspired Accent Stripe (Navy + Gold)
    bar_navy = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.06))
    bar_navy.fill.solid()
    bar_navy.fill.fore_color.rgb = NAVY_PRIMARY
    bar_navy.line.fill.background()
    
    bar_gold = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.06), Inches(2.2), Inches(0.035))
    bar_gold.fill.solid()
    bar_gold.fill.fore_color.rgb = GOLD_SOVEREIGN
    bar_gold.line.fill.background()
    
    # Kicker
    tb = s.shapes.add_textbox(Inches(0.8), Inches(0.38), Inches(9.5), Inches(0.28))
    tf = tb.text_frame
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = f"//  {kicker.upper()}"
    p.font.name = FONT_TECH
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = GOLD_SOVEREIGN
    
    # Title
    tb_t = s.shapes.add_textbox(Inches(0.8), Inches(0.66), Inches(11.0), Inches(0.65))
    tf_t = tb_t.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
    p2 = tf_t.paragraphs[0]
    p2.text = title
    p2.font.name = FONT_HEAD
    p2.font.size = Pt(21)
    p2.font.bold = True
    p2.font.color.rgb = NAVY_PRIMARY
    
    # Gold divider rule
    rule = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.36), Inches(1.8), Inches(0.04))
    rule.fill.solid()
    rule.fill.fore_color.rgb = GOLD_SOVEREIGN
    rule.line.fill.background()
    
    # Top Right Indicator
    tb_r = s.shapes.add_textbox(Inches(8.5), Inches(0.38), Inches(4.033), Inches(0.28))
    tf_r = tb_r.text_frame
    tf_r.margin_left = tf_r.margin_top = tf_r.margin_right = tf_r.margin_bottom = 0
    pr = tf_r.paragraphs[0]
    pr.alignment = PP_ALIGN.RIGHT
    pr.text = f"DEFENCE STRATEGIC BRIEFING  |  {slide_idx}"
    pr.font.name = FONT_TECH
    pr.font.size = Pt(8.5)
    pr.font.bold = True
    pr.font.color.rgb = TEXT_DIM
    
    # Footer
    tb_f = s.shapes.add_textbox(Inches(0.8), Inches(7.12), Inches(11.733), Inches(0.25))
    tf_f = tb_f.text_frame
    tf_f.margin_left = tf_f.margin_top = tf_f.margin_right = tf_f.margin_bottom = 0
    pf = tf_f.paragraphs[0]
    pf.text = "ATMANIRBHAR DEFENCE TECH  •  HEATBOLT SOVEREIGN THERMAL PLATFORM  •  RESTRICTED BRIEFING"
    pf.font.name = FONT_TECH
    pf.font.size = Pt(7.5)
    pf.font.color.rgb = TEXT_DIM
    
    return s

def add_card(s, x, y, w, h, title, desc="", tag=None, accent_color=GOLD_SOVEREIGN, title_size=11.5, desc_size=9, bg_col=PANEL_WHITE):
    c = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    c.fill.solid()
    c.fill.fore_color.rgb = bg_col
    c.line.color.rgb = BORDER_SUBTLE
    c.line.width = Pt(1)
    
    # Top accent line
    top_line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + 0.08), Inches(y), Inches(w - 0.16), Inches(0.035))
    top_line.fill.solid()
    top_line.fill.fore_color.rgb = accent_color
    top_line.line.fill.background()
    
    tb = s.shapes.add_textbox(Inches(x + 0.16), Inches(y + 0.12), Inches(w - 0.32), Inches(h - 0.22))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    if tag:
        ptag = tf.paragraphs[0]
        ptag.text = tag.upper()
        ptag.font.name = FONT_TECH
        ptag.font.size = Pt(8)
        ptag.font.bold = True
        ptag.font.color.rgb = accent_color
        ptag.space_after = Pt(2)
        ptitle = tf.add_paragraph()
    else:
        ptitle = tf.paragraphs[0]
        
    ptitle.text = title
    ptitle.font.name = FONT_HEAD
    ptitle.font.size = Pt(title_size)
    ptitle.font.bold = True
    ptitle.font.color.rgb = NAVY_PRIMARY
    ptitle.space_after = Pt(3)
    
    if desc:
        pdesc = tf.add_paragraph()
        pdesc.text = desc
        pdesc.font.name = FONT_BODY
        pdesc.font.size = Pt(desc_size)
        pdesc.font.color.rgb = TEXT_MUTED
        pdesc.space_before = Pt(2)

def add_tag_pills(s, x_start, y, tag_list, max_w=11.7):
    cur_x = x_start
    cur_y = y
    for tag in tag_list:
        w = max(0.9, len(tag) * 0.085 + 0.28)
        if cur_x + w > x_start + max_w:
            cur_x = x_start
            cur_y += 0.36
            
        pill = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cur_x), Inches(cur_y), Inches(w), Inches(0.28))
        pill.fill.solid()
        pill.fill.fore_color.rgb = TAG_BG_LIGHT
        pill.line.color.rgb = BORDER_SUBTLE
        pill.line.width = Pt(0.75)
        
        tb = s.shapes.add_textbox(Inches(cur_x), Inches(cur_y + 0.04), Inches(w), Inches(0.2))
        tf = tb.text_frame
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = tag
        p.font.name = FONT_TECH
        p.font.size = Pt(8)
        p.font.bold = True
        p.font.color.rgb = NAVY_SECONDARY
        
        cur_x += w + 0.12
    return cur_y + 0.36

def add_note_box(s, x, y, w, h, bold_prefix, note_text, accent=GOLD_SOVEREIGN):
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = GOLD_BG_TINT
    box.line.color.rgb = GOLD_BORDER
    box.line.width = Pt(1)
    
    # Left vertical accent
    left_bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y + 0.06), Inches(0.05), Inches(h - 0.12))
    left_bar.fill.solid()
    left_bar.fill.fore_color.rgb = accent
    left_bar.line.fill.background()
    
    tb = s.shapes.add_textbox(Inches(x + 0.18), Inches(y + 0.1), Inches(w - 0.3), Inches(h - 0.2))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    
    r1 = p.add_run()
    r1.text = bold_prefix + "  "
    r1.font.name = FONT_HEAD
    r1.font.bold = True
    r1.font.size = Pt(9.5)
    r1.font.color.rgb = accent
    
    r2 = p.add_run()
    r2.text = note_text
    r2.font.name = FONT_BODY
    r2.font.size = Pt(9)
    r2.font.color.rgb = TEXT_DARK

print("Slide builder helpers created.")
